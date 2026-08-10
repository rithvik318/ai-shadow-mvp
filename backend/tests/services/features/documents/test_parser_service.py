import io

import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.core.exceptions import (
    DocumentParseError,
    EmptyDocumentError,
    UnsupportedDocumentTypeError,
)
from app.services.features.documents.parser_service import (
    MAX_SECTION_TITLE_LENGTH,
    parse_document,
    resolve_format,
)
from tests.fixtures.factories import (
    build_docx,
    build_docx_alternate_content_text_box,
    build_docx_blocks,
    build_docx_nested_table,
    build_docx_table_with_text_box,
    build_docx_table_with_wrapped_runs,
    build_docx_wrapped_runs,
    build_markdown,
    build_pdf,
    build_pptx,
    build_pptx_merged_table,
    build_text,
)

PDF_TYPE = "application/pdf"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


# --- format resolution ---------------------------------------------------


@pytest.mark.parametrize(
    "filename, content_type, expected",
    [
        ("report.pdf", PDF_TYPE, "pdf"),
        ("report.docx", DOCX_TYPE, "docx"),
        ("deck.pptx", PPTX_TYPE, "pptx"),
        ("notes.txt", "text/plain", "txt"),
        ("notes.md", "text/markdown", "md"),
        ("notes.markdown", "text/x-markdown", "md"),
    ],
)
def test_resolve_format_uses_declared_content_type(
    filename: str, content_type: str, expected: str
) -> None:
    """A recognised content type determines the parser."""

    assert resolve_format(filename, content_type) == expected


@pytest.mark.parametrize(
    "filename, expected",
    [
        ("report.pdf", "pdf"),
        ("report.PDF", "pdf"),
        ("report.docx", "docx"),
        ("deck.pptx", "pptx"),
        ("deck.PPTX", "pptx"),
        ("notes.txt", "txt"),
        ("notes.md", "md"),
    ],
)
def test_resolve_format_falls_back_to_extension(filename: str, expected: str) -> None:
    """Clients sending a generic content type still resolve by extension."""

    assert resolve_format(filename, "application/octet-stream") == expected


def test_resolve_format_ignores_content_type_parameters() -> None:
    """A charset parameter does not defeat content-type matching."""

    assert resolve_format("notes.txt", "text/plain; charset=utf-8") == "txt"


@pytest.mark.parametrize(
    "filename, content_type",
    [
        ("archive.zip", "application/zip"),
        ("image.png", "image/png"),
        ("sheet.xlsx", "application/vnd.ms-excel"),
        ("noextension", None),
    ],
)
def test_resolve_format_rejects_unsupported_types(
    filename: str, content_type: str | None
) -> None:
    """Unsupported formats raise rather than being parsed as text."""

    with pytest.raises(UnsupportedDocumentTypeError):
        resolve_format(filename, content_type)


# --- PDF -----------------------------------------------------------------


def test_parse_pdf_returns_one_section_per_page_with_page_numbers() -> None:
    """PDF text is extracted per page, with 1-indexed page numbers."""

    pdf = build_pdf(
        ["Alpha page one content", "Beta page two content", "Gamma page three"]
    )

    parsed = parse_document(pdf, "report.pdf", PDF_TYPE)

    assert parsed.page_count == 3
    assert [section.page_number for section in parsed.sections] == [1, 2, 3]
    assert "Alpha" in parsed.sections[0].text
    assert "Gamma" in parsed.sections[2].text


def test_parse_pdf_skips_pages_with_no_extractable_text() -> None:
    """Blank pages are omitted, but still counted in page_count."""

    pdf = build_pdf(["Only this page has words", "", "Also this one"])

    parsed = parse_document(pdf, "report.pdf", PDF_TYPE)

    assert parsed.page_count == 3
    assert [section.page_number for section in parsed.sections] == [1, 3]


def test_parse_pdf_raises_when_no_page_has_text() -> None:
    """A PDF with no text layer is rejected rather than silently ingested."""

    pdf = build_pdf(["", ""])

    with pytest.raises(EmptyDocumentError):
        parse_document(pdf, "scanned.pdf", PDF_TYPE)


def test_parse_pdf_raises_document_parse_error_for_corrupt_bytes() -> None:
    """Bytes claiming to be a PDF but which are not raise a parse error."""

    with pytest.raises(DocumentParseError):
        parse_document(b"this is definitely not a pdf", "broken.pdf", PDF_TYPE)


# --- DOCX ----------------------------------------------------------------


def test_parse_docx_groups_paragraphs_under_their_heading() -> None:
    """Each heading starts a new section carrying that heading as its title."""

    data = build_docx(
        [
            ("Introduction", "This is the introduction body."),
            ("Methods", "This is the methods body."),
        ]
    )

    parsed = parse_document(data, "paper.docx", DOCX_TYPE)

    assert [section.section_title for section in parsed.sections] == [
        "Introduction",
        "Methods",
    ]
    assert "introduction body" in parsed.sections[0].text
    assert parsed.page_count is None


def test_parse_docx_reports_no_page_numbers() -> None:
    """DOCX has no reliable pagination, so page_number stays unset."""

    data = build_docx([("Title", "Body")])

    parsed = parse_document(data, "paper.docx", DOCX_TYPE)

    assert all(section.page_number is None for section in parsed.sections)


def test_parse_docx_raises_document_parse_error_for_corrupt_bytes() -> None:
    """Unreadable DOCX bytes surface as a parse error."""

    with pytest.raises(DocumentParseError):
        parse_document(b"not a real docx", "broken.docx", DOCX_TYPE)


# --- DOCX tables and text boxes ------------------------------------------


def test_parse_docx_with_only_paragraphs_is_unchanged_by_block_walking() -> None:
    """Walking the body must not alter output for a plain prose document.

    Guards the backward-compatibility property the rest of the DOCX tests
    assume: a file with no tables and no shapes parses exactly as before.
    """

    blocks = build_docx_blocks(
        [
            ("heading", "Introduction"),
            ("paragraph", "First body paragraph."),
            ("paragraph", "Second body paragraph."),
        ]
    )

    parsed = parse_document(blocks, "prose.docx", DOCX_TYPE)

    assert [(s.section_title, s.text) for s in parsed.sections] == [
        ("Introduction", "First body paragraph.\nSecond body paragraph.")
    ]


def test_parse_docx_extracts_a_document_that_is_only_a_table() -> None:
    """A file whose entire content is tabular is no longer empty."""

    data = build_docx_blocks(
        [("table", [["Metric", "Value"], ["Revenue", "10"], ["Margin", "22%"]])]
    )

    parsed = parse_document(data, "figures.docx", DOCX_TYPE)

    assert [section.text for section in parsed.sections] == [
        "Metric: Revenue | Value: 10\nMetric: Margin | Value: 22%"
    ]


def test_parse_docx_extracts_a_document_that_is_only_text_boxes() -> None:
    """Shape-only layouts extracted as empty before; they now yield text."""

    data = build_docx_blocks([("textbox", "Core capability\nData governance")])

    parsed = parse_document(data, "capability.docx", DOCX_TYPE)

    assert [section.text for section in parsed.sections] == [
        "Core capability\nData governance"
    ]


def test_parse_docx_keeps_paragraphs_and_tables_in_document_order() -> None:
    """A table between two paragraphs stays between them.

    `document.paragraphs` and `document.tables` are separate flat lists, so
    reading them independently would move every table to the end.
    """

    data = build_docx_blocks(
        [
            ("paragraph", "Before the table."),
            ("table", [["Name", "Role"], ["Ann", "Lead"]]),
            ("paragraph", "After the table."),
        ]
    )

    parsed = parse_document(data, "ordered.docx", DOCX_TYPE)

    assert [section.text for section in parsed.sections] == [
        "Before the table.",
        "Name: Ann | Role: Lead",
        "After the table.",
    ]


def test_parse_docx_keeps_mixed_content_in_document_order() -> None:
    """Headings, prose, tables and shapes interleave in body order."""

    data = build_docx_blocks(
        [
            ("heading", "Overview"),
            ("paragraph", "Opening prose."),
            ("table", [["Name", "Role"], ["Ann", "Lead"]]),
            ("textbox", "Pull quote"),
            ("heading", "Detail"),
            ("paragraph", "Closing prose."),
        ]
    )

    parsed = parse_document(data, "mixed.docx", DOCX_TYPE)

    assert [(s.section_title, s.text) for s in parsed.sections] == [
        ("Overview", "Opening prose."),
        ("Overview", "Name: Ann | Role: Lead"),
        ("Overview", "Pull quote"),
        ("Detail", "Closing prose."),
    ]


def test_parse_docx_qualifies_each_table_row_with_its_header() -> None:
    """Every data row restates its column labels.

    A chunk boundary can fall between any two rows, so a row that has been
    separated from a header line would otherwise lose the meaning of its
    values.
    """

    data = build_docx_blocks(
        [
            (
                "table",
                [
                    ["Contract", "Agency", "Value"],
                    ["SF-DART", "DOT", "1.2M"],
                    ["SF-RAIL", "FTA", "800K"],
                ],
            )
        ]
    )

    parsed = parse_document(data, "contracts.docx", DOCX_TYPE)

    assert parsed.sections[0].text == (
        "Contract: SF-DART | Agency: DOT | Value: 1.2M\n"
        "Contract: SF-RAIL | Agency: FTA | Value: 800K"
    )


def test_parse_docx_keeps_a_table_out_of_the_surrounding_prose_section() -> None:
    """Table text becomes its own section so a chunk never straddles both."""

    data = build_docx_blocks(
        [
            ("paragraph", "Prose."),
            ("table", [["A", "B"], ["1", "2"]]),
        ]
    )

    parsed = parse_document(data, "split.docx", DOCX_TYPE)

    assert len(parsed.sections) == 2
    assert "A: 1" not in parsed.sections[0].text


def test_parse_docx_collapses_a_horizontally_merged_cell() -> None:
    """A merged cell is reported once, not once per column it spans."""

    data = build_docx_blocks([("merged", [["Name", "Role"], ["Ann", "Lead"]])])

    parsed = parse_document(data, "merged.docx", DOCX_TYPE)

    text = parsed.sections[0].text
    assert text.count("Name Role") == 1
    assert text.endswith("Ann | Lead")


def test_parse_docx_skips_entirely_empty_table_rows() -> None:
    """Spacer rows contribute no line rather than an empty one."""

    data = build_docx_blocks(
        [("table", [["Name", "Role"], ["", ""], ["Ann", "Lead"], ["", ""]])]
    )

    parsed = parse_document(data, "spaced.docx", DOCX_TYPE)

    assert parsed.sections[0].text == "Name: Ann | Role: Lead"


def test_parse_docx_falls_back_to_plain_rows_for_a_single_column_table() -> None:
    """A one-column layout table has no labels, so nothing is prefixed."""

    data = build_docx_blocks([("table", [["First line"], ["Second line"]])])

    parsed = parse_document(data, "layout.docx", DOCX_TYPE)

    assert parsed.sections[0].text == "First line\nSecond line"


def test_parse_docx_falls_back_to_plain_rows_when_headers_repeat() -> None:
    """Repeated first-row values are formatting, not column labels."""

    data = build_docx_blocks([("table", [["Cell", "Cell"], ["Left", "Right"]])])

    parsed = parse_document(data, "grid.docx", DOCX_TYPE)

    assert parsed.sections[0].text == "Cell | Cell\nLeft | Right"


def test_parse_docx_keeps_a_header_only_table() -> None:
    """A single-row table still contributes its text."""

    data = build_docx_blocks([("table", [["Alpha", "Beta"]])])

    parsed = parse_document(data, "header.docx", DOCX_TYPE)

    assert parsed.sections[0].text == "Alpha | Beta"


def test_parse_docx_reads_a_table_nested_inside_a_cell() -> None:
    """`_Cell.text` sees only direct paragraphs, so nesting is walked."""

    data = build_docx_nested_table("Outer label", [["k1", "v1"], ["k2", "v2"]])

    parsed = parse_document(data, "nested.docx", DOCX_TYPE)

    text = parsed.sections[0].text
    assert "Outer label" in text
    assert "k2" in text and "v2" in text


def test_parse_docx_reads_a_text_box_inside_a_table_cell() -> None:
    """Shapes anchored inside cells are part of that cell's text."""

    data = build_docx_table_with_text_box(
        [["Name", "Note"], ["Ann", "Plain"]], "Shape note"
    )

    parsed = parse_document(data, "cellbox.docx", DOCX_TYPE)

    assert "Shape note" in parsed.sections[0].text


def test_parse_docx_places_a_text_box_after_its_anchoring_paragraph() -> None:
    """A shape reads after the paragraph it is anchored to, not before."""

    data = build_docx_blocks([("paragraph", "Anchor text."), ("textbox", "Shape text")])

    parsed = parse_document(data, "anchor.docx", DOCX_TYPE)

    assert [section.text for section in parsed.sections] == [
        "Anchor text.",
        "Shape text",
    ]


def test_parse_docx_reads_a_dual_encoded_text_box_once() -> None:
    """Word's `mc:Choice` / `mc:Fallback` pair is one shape, not two.

    Regression: reading both branches doubled every shape in the corpus's
    capability statements.
    """

    data = build_docx_alternate_content_text_box("Core Competencies")

    parsed = parse_document(data, "shape.docx", DOCX_TYPE)

    assert [section.text for section in parsed.sections] == ["Core Competencies"]


def test_parse_docx_rejects_a_document_with_no_content() -> None:
    """An empty DOCX yields no sections and is rejected, not stored blank."""

    with pytest.raises(EmptyDocumentError):
        parse_document(build_docx_blocks([]), "empty.docx", DOCX_TYPE)


def test_parse_docx_rejects_a_document_of_only_empty_tables() -> None:
    """A table of blank cells is no more ingestible than a blank file."""

    data = build_docx_blocks([("table", [["", ""], ["", ""]])])

    with pytest.raises(EmptyDocumentError):
        parse_document(data, "blank-table.docx", DOCX_TYPE)


# --- DOCX headings with no body beneath them -----------------------------


def test_parse_docx_keeps_a_heading_superseded_by_another_heading() -> None:
    """A heading only ever survived as the title of a later section.

    Two headings in a row overwrote the first before anything carried it, and
    nothing failed to show for it.
    """

    data = build_docx_blocks(
        [("heading", "Heading A"), ("heading", "Heading B"), ("paragraph", "Body B")]
    )

    parsed = parse_document(data, "headings.docx", DOCX_TYPE)

    assert [(s.section_title, s.text) for s in parsed.sections] == [
        ("Heading A", "Heading A"),
        ("Heading B", "Body B"),
    ]


def test_parse_docx_keeps_every_heading_in_a_run_of_three() -> None:
    """Each superseded heading becomes its own section, in document order."""

    data = build_docx_blocks(
        [
            ("heading", "Heading A"),
            ("heading", "Heading B"),
            ("heading", "Heading C"),
            ("paragraph", "Body C"),
        ]
    )

    parsed = parse_document(data, "headings.docx", DOCX_TYPE)

    assert [s.text for s in parsed.sections] == ["Heading A", "Heading B", "Body C"]


def test_parse_docx_accepts_a_document_that_is_only_a_heading() -> None:
    """A heading with nothing after it used to yield no sections at all.

    The document was then rejected as empty, so a file full of visible text
    was recorded as failed with no explanation.
    """

    parsed = parse_document(
        build_docx_blocks([("heading", "Statement of Work")]), "h.docx", DOCX_TYPE
    )

    assert [(s.section_title, s.text) for s in parsed.sections] == [
        ("Statement of Work", "Statement of Work")
    ]


def test_parse_docx_keeps_a_trailing_heading() -> None:
    """A heading ending the document is preserved after the body above it."""

    data = build_docx_blocks([("paragraph", "Body A"), ("heading", "Heading B")])

    parsed = parse_document(data, "trailing.docx", DOCX_TYPE)

    assert [(s.section_title, s.text) for s in parsed.sections] == [
        (None, "Body A"),
        ("Heading B", "Heading B"),
    ]


def test_parse_docx_does_not_repeat_a_heading_that_titles_a_table() -> None:
    """A table already carries the heading above it, so it is not restated."""

    data = build_docx_blocks(
        [
            ("heading", "Heading A"),
            ("table", [["K", "V"], ["k1", "v1"]]),
            ("heading", "Heading B"),
        ]
    )

    parsed = parse_document(data, "tabled.docx", DOCX_TYPE)

    assert [(s.section_title, s.text) for s in parsed.sections] == [
        ("Heading A", "K: k1 | V: v1"),
        ("Heading B", "Heading B"),
    ]


def test_parse_docx_does_not_repeat_a_heading_that_titles_a_text_box() -> None:
    """The same holds for a shape: it carries the heading, so it is not repeated."""

    data = build_docx_blocks(
        [
            ("heading", "Heading A"),
            ("textbox", "Boxed text"),
            ("heading", "Heading B"),
        ]
    )

    parsed = parse_document(data, "boxed.docx", DOCX_TYPE)

    assert [(s.section_title, s.text) for s in parsed.sections] == [
        ("Heading A", "Boxed text"),
        ("Heading B", "Heading B"),
    ]


def test_parse_docx_alternating_headings_and_bodies_are_unchanged() -> None:
    """The control: the ordinary shape must parse exactly as it always did."""

    data = build_docx([("Introduction", "Intro body."), ("Methods", "Methods body.")])

    parsed = parse_document(data, "paper.docx", DOCX_TYPE)

    assert [(s.section_title, s.text) for s in parsed.sections] == [
        ("Introduction", "Intro body."),
        ("Methods", "Methods body."),
    ]


def test_parse_docx_never_produces_a_section_with_empty_text() -> None:
    """Recovering a heading must not cost a crop of empty sections."""

    data = build_docx_blocks(
        [
            ("heading", "A"),
            ("heading", "B"),
            ("paragraph", "Body"),
            ("table", [["K", "V"], ["k", "v"]]),
            ("textbox", "Boxed"),
            ("heading", "C"),
        ]
    )

    parsed = parse_document(data, "mixed.docx", DOCX_TYPE)

    assert parsed.sections
    assert all(section.text.strip() for section in parsed.sections)


# --- section_title is bounded, the heading is not ------------------------


def test_parse_docx_leaves_a_title_at_the_limit_untouched() -> None:
    """A heading exactly at the bound is not trimmed."""

    heading = "H" * MAX_SECTION_TITLE_LENGTH
    data = build_docx_blocks([("heading", heading), ("paragraph", "Body.")])

    parsed = parse_document(data, "limit.docx", DOCX_TYPE)

    assert parsed.sections[0].section_title == heading


def test_parse_docx_bounds_a_long_title_but_keeps_the_whole_heading() -> None:
    """The label is trimmed; the heading itself is not.

    `section_title` is metadata — never embedded, never searched — so bounding
    it loses nothing retrievable, provided the heading survives as text.
    """

    heading = "Solution Approach: " + "word " * 200
    data = build_docx_blocks([("heading", heading.strip())])

    parsed = parse_document(data, "long.docx", DOCX_TYPE)
    section = parsed.sections[0]

    assert len(section.section_title) == MAX_SECTION_TITLE_LENGTH
    assert section.text == heading.strip()
    assert len(section.text) > MAX_SECTION_TITLE_LENGTH


def test_parse_docx_handles_the_corpus_636_character_heading() -> None:
    """The real case: a 636-character paragraph styled as a heading.

    `DocumentChunk.section_title` is `String(512)`, which PostgreSQL enforces
    and SQLite does not — so before this bound the document parsed cleanly in
    every test and failed the insert on the database that ships.
    """

    # Exactly 636 characters and ending on a word, so the parser's `strip()`
    # cannot make the assertion pass or fail for the wrong reason.
    sentence = "For over 15 years Aikya has extensively provided data services. "
    heading = (sentence * 10)[:635] + "."
    assert len(heading) == 636

    data = build_docx_blocks([("heading", heading), ("paragraph", "Body.")])

    parsed = parse_document(data, "exhibit.docx", DOCX_TYPE)

    assert len(parsed.sections[0].section_title) <= MAX_SECTION_TITLE_LENGTH
    assert heading in parsed.text


def test_parse_pptx_bounds_a_long_slide_title() -> None:
    """PPTX titles go through the same bound, and stay whole in the text."""

    title = "Product Information Management Concepts and Their Application " * 6
    title = title.strip()
    parsed = parse_document(build_pptx([{"title": title}]), "deck.pptx", PPTX_TYPE)
    section = parsed.sections[0]

    assert len(title) > MAX_SECTION_TITLE_LENGTH
    assert len(section.section_title) == MAX_SECTION_TITLE_LENGTH
    assert section.text == title


@pytest.mark.parametrize("length", [10, 199, 200, 201, 636, 2000])
def test_every_section_title_respects_the_bound(length: int) -> None:
    """The invariant, across both formats that produce titles."""

    heading = "T" * length
    docx_parsed = parse_document(
        build_docx_blocks([("heading", heading), ("paragraph", "Body.")]),
        "d.docx",
        DOCX_TYPE,
    )
    pptx_parsed = parse_document(
        build_pptx([{"title": heading, "body": "Body."}]), "d.pptx", PPTX_TYPE
    )

    for parsed in (docx_parsed, pptx_parsed):
        for section in parsed.sections:
            if section.section_title is not None:
                assert len(section.section_title) <= MAX_SECTION_TITLE_LENGTH


# --- DOCX runs nested inside wrappers ------------------------------------


@pytest.mark.parametrize("wrapper", ["ins", "hyperlink", "fldSimple", "smartTag"])
def test_parse_docx_reads_a_run_wrapped_in_a_container(wrapper: str) -> None:
    """Text is extracted whichever wrapper Word puts between it and the run.

    `Paragraph.text` reads direct runs and hyperlinks only, so a tracked
    insertion, a field result or a smart tag was previously invisible.
    """

    data = build_docx_wrapped_runs([[(wrapper, "Load-bearing sentence.")]])

    parsed = parse_document(data, "wrapped.docx", DOCX_TYPE)

    assert parsed.sections[0].text == "Load-bearing sentence."


def test_parse_docx_joins_direct_and_nested_runs_in_one_paragraph() -> None:
    """Runs concatenate in document order, whatever wraps each one.

    Joined without a separator because Word splits a single word across runs
    freely; anything else would insert a break mid-word.
    """

    data = build_docx_wrapped_runs(
        [
            [
                ("direct", "Data "),
                ("ins", "governance "),
                ("smartTag", "Reston "),
                ("fldSimple", "42 "),
                ("hyperlink", "www.SunRadia.com"),
            ]
        ]
    )

    parsed = parse_document(data, "mixed.docx", DOCX_TYPE)

    assert parsed.sections[0].text == "Data governance Reston 42 www.SunRadia.com"


def test_parse_docx_excludes_deleted_text() -> None:
    """A tracked deletion is not content and must not be indexed."""

    data = build_docx_wrapped_runs(
        [[("direct", "Kept. "), ("del", "Struck out. "), ("ins", "Added.")]]
    )

    parsed = parse_document(data, "tracked.docx", DOCX_TYPE)

    assert parsed.sections[0].text == "Kept. Added."
    assert "Struck out" not in parsed.text


def test_parse_docx_keeps_paragraph_boundaries_across_wrappers() -> None:
    """Each paragraph stays its own line, and their order is preserved."""

    data = build_docx_wrapped_runs(
        [
            [("ins", "First paragraph.")],
            [("direct", "Second paragraph.")],
            [("fldSimple", "Third paragraph.")],
        ]
    )

    parsed = parse_document(data, "ordered.docx", DOCX_TYPE)

    assert parsed.sections[0].text == (
        "First paragraph.\nSecond paragraph.\nThird paragraph."
    )


def test_parse_docx_does_not_reject_a_document_whose_text_is_all_inserted() -> None:
    """The real-corpus failure: a document of unaccepted tracked changes.

    `LN_ECM_FDD Physical Data Model_DJP_04292009_v3.1.docx` has 570 paragraphs
    of which `Paragraph.text` saw 17. It parsed "successfully" and indexed as
    an empty shell, which is worse than failing.
    """

    data = build_docx_wrapped_runs(
        [[("ins", "Table of Contents")], [("ins", "Design Approach for the model")]]
    )

    parsed = parse_document(data, "tracked.docx", DOCX_TYPE)

    assert parsed.sections[0].text == (
        "Table of Contents\nDesign Approach for the model"
    )


def test_parse_docx_reads_a_wrapped_run_inside_a_table_cell() -> None:
    """Cell text goes through the same paragraph reader as body text."""

    data = build_docx_table_with_wrapped_runs("Lead architect")

    parsed = parse_document(data, "celltracked.docx", DOCX_TYPE)

    assert parsed.sections[0].text == "Name: Ann | Role: Lead architect"


def test_parse_docx_counts_a_wrapped_run_as_a_heading() -> None:
    """A heading whose text is inserted still titles its section."""

    document = docx.Document()
    heading = document.add_heading("", level=1)
    heading._p.append(
        parse_xml(
            f'<w:ins {nsdecls("w")} w:id="9" w:author="A"'
            f' w:date="2024-01-01T00:00:00Z">'
            f"<w:r><w:t>Executive Summary</w:t></w:r></w:ins>"
        )
    )
    document.add_paragraph("Body text.")
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = parse_document(buffer.getvalue(), "heading.docx", DOCX_TYPE)

    assert parsed.sections[0].section_title == "Executive Summary"
    assert parsed.sections[0].text == "Body text."


# --- PPTX ----------------------------------------------------------------


def test_parse_pptx_returns_one_section_per_slide_with_slide_numbers() -> None:
    """Each slide becomes a section, numbered from one, in slide order."""

    data = build_pptx(
        [{"title": "One"}, {"title": "Two"}, {"title": "Three"}],
    )

    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert parsed.page_count == 3
    assert [section.page_number for section in parsed.sections] == [1, 2, 3]
    assert [section.text for section in parsed.sections] == ["One", "Two", "Three"]


def test_parse_pptx_extracts_a_title_slide() -> None:
    """A title placeholder becomes both the section title and its first line.

    Repeated rather than only stored as `section_title` because a divider slide
    often has no other content, and dropping it would drop the slide.
    """

    parsed = parse_document(
        build_pptx([{"title": "Data Governance"}]), "d.pptx", PPTX_TYPE
    )

    assert parsed.sections[0].section_title == "Data Governance"
    assert parsed.sections[0].text == "Data Governance"


def test_parse_pptx_extracts_a_body_placeholder_under_its_title() -> None:
    """Bullet text follows the title it belongs to."""

    data = build_pptx([{"title": "Approach", "body": "Assess\nDesign\nDeliver"}])

    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert parsed.sections[0].text == "Approach\nAssess\nDesign\nDeliver"


def test_parse_pptx_orders_shapes_by_position_not_by_z_order() -> None:
    """Reading order follows the slide, not the order shapes were added.

    The fixture adds the lowest box first, so a parser that trusted the shape
    collection's own order would emit the slide bottom-up.
    """

    data = build_pptx([{"boxes": [(5, "Bottom"), (1, "Top"), (3, "Middle")]}])

    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert parsed.sections[0].text == "Top\nMiddle\nBottom"


def test_parse_pptx_keeps_the_title_first_regardless_of_its_position() -> None:
    """The title leads even when a shape sits above it on the slide."""

    data = build_pptx([{"title": "Heading", "boxes": [(0.1, "Sits above the title")]}])

    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert parsed.sections[0].text.splitlines()[0] == "Heading"


def test_parse_pptx_reads_shapes_inside_a_group() -> None:
    """Grouped shapes are walked through, in their own reading order."""

    data = build_pptx(
        [{"boxes": [(1, "Before")], "group": [(5, "Group lower"), (3, "Group upper")]}]
    )

    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert parsed.sections[0].text == "Before\nGroup upper\nGroup lower"


def test_parse_pptx_serialises_a_table_the_same_way_docx_does() -> None:
    """One `Header: value` line per row, identical to the DOCX serialiser."""

    grid = [["Metric", "Value"], ["Revenue", "10"], ["Margin", "22%"]]

    from_pptx = parse_document(build_pptx([{"table": grid}]), "d.pptx", PPTX_TYPE)
    from_docx = parse_document(
        build_docx_blocks([("table", grid)]), "d.docx", DOCX_TYPE
    )

    assert from_pptx.sections[0].text == from_docx.sections[0].text
    assert from_pptx.sections[0].text == (
        "Metric: Revenue | Value: 10\nMetric: Margin | Value: 22%"
    )


def test_parse_pptx_collapses_a_merged_table_cell() -> None:
    """A spanned cell is reported once, not once per column it covers."""

    data = build_pptx_merged_table([["Name", "Role"], ["Ann", "Lead"]])

    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert parsed.sections[0].text.count("Name Role") == 1
    assert parsed.sections[0].text.endswith("Ann | Lead")


def test_parse_pptx_appends_labelled_speaker_notes() -> None:
    """Notes come last and are marked as the presenter's, not the slide's."""

    data = build_pptx([{"title": "Pricing", "notes": "Do not quote a number"}])

    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert parsed.sections[0].text == ("Pricing\nSpeaker notes: Do not quote a number")


def test_parse_pptx_handles_a_mixed_layout_slide() -> None:
    """Title, body, free text box, table and notes on one slide, in order."""

    data = build_pptx(
        [
            {
                "title": "Engagement",
                "body": "Scope agreed",
                "boxes": [(4, "Footnote box")],
                "table": [["Phase", "Weeks"], ["Discovery", "4"]],
                "notes": "Mention the pilot",
            }
        ]
    )

    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert parsed.sections[0].text == (
        "Engagement\n"
        "Scope agreed\n"
        "Footnote box\n"
        "Phase: Discovery | Weeks: 4\n"
        "Speaker notes: Mention the pilot"
    )


def test_parse_pptx_includes_hidden_slides() -> None:
    """Hidden slides are extracted, deliberately.

    A hidden slide is authored content that was set aside for one audience —
    usually backup detail worth answering questions from. Skipping it would be
    a silent content loss of exactly the kind this parser exists to stop.
    """

    data = build_pptx([{"title": "Shown"}, {"title": "Backup detail", "hidden": True}])

    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert [section.text for section in parsed.sections] == ["Shown", "Backup detail"]


def test_parse_pptx_skips_slides_with_no_text_but_still_counts_them() -> None:
    """An image-only slide contributes no section, as a blank PDF page does."""

    data = build_pptx([{"title": "Has text"}, {}, {"title": "Also has text"}])

    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert parsed.page_count == 3
    assert [section.page_number for section in parsed.sections] == [1, 3]


def test_parse_pptx_rejects_a_presentation_with_no_slides() -> None:
    """An empty deck is rejected rather than stored with nothing in it."""

    with pytest.raises(EmptyDocumentError):
        parse_document(build_pptx([]), "empty.pptx", PPTX_TYPE)


def test_parse_pptx_rejects_a_presentation_with_no_text() -> None:
    """Slides exist but none carries text, so there is nothing to index."""

    with pytest.raises(EmptyDocumentError):
        parse_document(build_pptx([{}, {}]), "blank.pptx", PPTX_TYPE)


def test_parse_pptx_raises_document_parse_error_for_corrupt_bytes() -> None:
    """Bytes claiming to be a PPTX but which are not raise a parse error."""

    with pytest.raises(DocumentParseError):
        parse_document(b"this is not a presentation", "broken.pptx", PPTX_TYPE)


def test_parse_pptx_recovers_from_a_slide_it_cannot_read(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """One unreadable slide costs its own content, not the whole deck."""

    from app.services.features.documents import parser_service

    original = parser_service._slide_blocks
    calls: list[int] = []

    def exploding(slide):
        calls.append(1)
        if len(calls) == 2:
            raise ValueError("malformed shape tree")
        return original(slide)

    monkeypatch.setattr(parser_service, "_slide_blocks", exploding)

    data = build_pptx([{"title": "One"}, {"title": "Two"}, {"title": "Three"}])
    parsed = parse_document(data, "deck.pptx", PPTX_TYPE)

    assert parsed.page_count == 3
    assert [section.text for section in parsed.sections] == ["One", "Three"]


# --- Markdown ------------------------------------------------------------


def test_parse_markdown_splits_on_headings() -> None:
    """ATX headings become section titles and start new sections."""

    data = build_markdown(
        [("Overview", "Overview body text."), ("Details", "Details body text.")]
    )

    parsed = parse_document(data, "notes.md", "text/markdown")

    assert [section.section_title for section in parsed.sections] == [
        "Overview",
        "Details",
    ]
    assert "Overview body text." in parsed.sections[0].text


def test_parse_markdown_keeps_content_before_the_first_heading() -> None:
    """A preamble with no heading is preserved with no section title."""

    data = b"Preamble paragraph.\n\n# First Heading\n\nSection body."

    parsed = parse_document(data, "notes.md", "text/markdown")

    assert parsed.sections[0].section_title is None
    assert "Preamble paragraph." in parsed.sections[0].text
    assert parsed.sections[1].section_title == "First Heading"


# --- Plain text ----------------------------------------------------------


def test_parse_text_returns_a_single_untitled_section() -> None:
    """Plain text has no structure to preserve."""

    parsed = parse_document(build_text("Line one\nLine two"), "n.txt", "text/plain")

    assert len(parsed.sections) == 1
    assert parsed.sections[0].section_title is None
    assert parsed.sections[0].page_number is None
    assert parsed.sections[0].text == "Line one\nLine two"


def test_parse_text_decodes_non_utf8_bytes() -> None:
    """Windows-encoded text is decoded rather than rejected."""

    parsed = parse_document("Café notes".encode("cp1252"), "notes.txt", "text/plain")

    assert "notes" in parsed.sections[0].text


def test_parse_text_strips_a_utf8_byte_order_mark() -> None:
    """A BOM does not leak into the extracted text."""

    parsed = parse_document(b"\xef\xbb\xbfHello", "notes.txt", "text/plain")

    assert parsed.sections[0].text == "Hello"


# --- Empty input ---------------------------------------------------------


def test_parse_document_rejects_empty_bytes() -> None:
    """A zero-byte upload is rejected before format resolution."""

    with pytest.raises(EmptyDocumentError):
        parse_document(b"", "notes.txt", "text/plain")


@pytest.mark.parametrize("body", ["", "   ", "\n\n\t "])
def test_parse_document_rejects_whitespace_only_text(body: str) -> None:
    """Whitespace-only content yields no sections and is rejected."""

    with pytest.raises(EmptyDocumentError):
        parse_document(build_text(body) or b" ", "notes.txt", "text/plain")


def test_parsed_document_text_joins_all_sections() -> None:
    """The convenience `text` property concatenates sections in order."""

    parsed = parse_document(
        build_markdown([("A", "first"), ("B", "second")]),
        "notes.md",
        "text/markdown",
    )

    assert parsed.text == "first\n\nsecond"
